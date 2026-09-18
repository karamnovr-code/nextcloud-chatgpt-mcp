"""Behavioral acceptance fixtures for the isolated document engine."""
from io import BytesIO
import json
import zipfile

import fitz
import pytest
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from openpyxl import Workbook, load_workbook

import doc_engine as engine


def pdf_fixture(scanned=False):
    doc = fitz.open()
    for number in range(1, 4):
        page = doc.new_page()
        if scanned:
            im = Image.new('RGB', (1500, 600), 'white')
            draw = ImageDraw.Draw(im)
            font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 52)
            draw.text((60, 100), f'Документ номер {number}', fill='black', font=font)
            draw.text((60, 200), 'Проверка русского текста 12345', fill='black', font=font)
            stream = BytesIO(); im.save(stream, 'PNG')
            page.insert_image(fitz.Rect(0, 0, 595, 238), stream=stream.getvalue())
        else:
            page.insert_text((60, 100), f'Page {number}: verified text 12345')
    result = doc.tobytes(); doc.close()
    return result


def test_pdf_pagination_has_exact_page_sources_and_no_silent_truncation():
    data = pdf_fixture()
    first = engine.read_document(data, 'report.pdf', start=1, count=2)
    assert first['total_units'] == 3
    assert first['next_start'] == 3
    assert [u['page'] for u in first['units']] == [1, 2]
    assert 'Page 2' in first['units'][1]['text']
    assert first['units'][1]['source'] == 'report.pdf#page=2'
    assert engine.read_document(data, 'report.pdf', start=3)['next_start'] is None
    json.dumps(first, ensure_ascii=False)


def test_scanned_russian_pdf_is_read_page_by_page():
    result = engine.read_document(pdf_fixture(True), 'scan.pdf', start=2, count=1)
    assert result['total_units'] == 3
    unit = result['units'][0]
    assert unit['page'] == 2
    assert 'Документ номер 2' in unit['text']
    assert 'русского текста 12345' in unit['text']
    assert unit['ocr']['attempted'] is True
    assert unit['ocr']['languages'] == 'rus+eng'
    assert 'incomplete' in unit['ocr']
    assert result['next_start'] == 3


def test_pdf_preview_is_actual_requested_page_png():
    data = engine.preview(pdf_fixture(), 'report.pdf', page=2)
    assert data.startswith(b'\x89PNG\r\n\x1a\n')
    im = Image.open(BytesIO(data))
    assert im.width >= 595 and im.height >= 842
    with pytest.raises(ValueError):
        engine.preview(pdf_fixture(), 'report.pdf', page=0)


def test_docx_create_and_template_replacement_preserve_runs_and_source():
    template = Document()
    p = template.add_paragraph()
    p.add_run('Dear {{na').bold = True
    p.add_run('me}}, welcome.').italic = True
    table = template.add_table(rows=1, cols=1)
    table.cell(0, 0).text = '{{name}}'
    before = BytesIO(); template.save(before)
    raw = before.getvalue()
    output = engine.create_docx({'replacements': {'{{name}}': 'Роман'}, 'paragraphs': ['Проверено']}, raw)
    doc = Document(BytesIO(output))
    assert doc.paragraphs[0].text == 'Dear Роман, welcome.'
    assert doc.paragraphs[0].runs[0].bold
    assert doc.tables[0].cell(0, 0).text == 'Роман'
    assert Document(BytesIO(raw)).paragraphs[0].text == 'Dear {{name}}, welcome.'
    read = engine.read_document(output, 'letter.docx', count=50)
    assert any('Проверено' in unit['text'] for unit in read['units'])
    assert read['unit_type'] == 'block'


def test_xlsx_literals_formulas_copy_update_and_readback():
    original = engine.create_xlsx({'sheets': [{'name': 'Расчёт', 'rows': [['Описание', 'Значение'], ['=1+1', 5], ['Итого', {'formula': '=SUM(B2:B2)'}]]}]})
    wb = load_workbook(BytesIO(original))
    assert wb['Расчёт']['A2'].value == '=1+1'
    assert wb['Расчёт']['A2'].data_type == 's'
    assert wb['Расчёт']['B3'].data_type == 'f'
    updated = engine.create_xlsx({'sheets': [{'name': 'Расчёт', 'cells': {'B2': 7}}]}, original)
    assert load_workbook(BytesIO(updated))['Расчёт']['B2'].value == 7
    assert load_workbook(BytesIO(original))['Расчёт']['B2'].value == 5
    result = engine.read_document(updated, 'budget.xlsx', count=20)
    assert result['total_units'] == 3
    formula = result['units'][2]['cells'][1]
    assert formula['formula'] == '=SUM(B2:B2)'
    assert formula['value'] is None
    assert formula['cached_value_available'] is False
    assert result['units'][2]['source'] == 'budget.xlsx#sheet=Расчёт&row=3'


def test_image_exif_and_ocr_are_json_safe():
    im = Image.new('RGB', (500, 200), 'white')
    exif = Image.Exif(); exif[271] = 'Fixture camera'
    out = BytesIO(); im.save(out, 'JPEG', exif=exif)
    result = engine.read_document(out.getvalue(), 'photo.jpg')
    assert result['metadata']['exif']['Make'] == 'Fixture camera'
    assert result['metadata']['width'] == 500
    json.dumps(result)
    assert engine.preview(out.getvalue(), 'photo.jpg').startswith(b'\x89PNG')


@pytest.mark.parametrize('filename', ['../../secret.pdf', '/tmp/secret.pdf', 'evil.xlsm'])
def test_reject_unsafe_filename_or_active_office_format(filename):
    with pytest.raises(ValueError):
        engine.read_document(b'abc', filename)


def test_reject_zip_traversal_and_expansion_bomb():
    out = BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('../word/document.xml', '<x/>')
    with pytest.raises(ValueError):
        engine.read_document(out.getvalue(), 'evil.docx')
    out = BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('word/document.xml', b'0' * (16 * 1024 * 1024))
    with pytest.raises(ValueError, match='(?i)(compression|expansion|archive)'):
        engine.read_document(out.getvalue(), 'bomb.docx')


def test_pdf_conversion_preserves_all_pages_with_caveat():
    raw = engine.convert_document(pdf_fixture(), 'original.pdf', 'docx')
    doc = Document(BytesIO(raw))
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert 'Page 1' in text and 'Page 3' in text
    assert 'best effort' in doc.core_properties.comments.lower()


def test_office_pdf_conversion_renders_russian_text_and_preview():
    raw = engine.create_docx({'title': 'Проверка документа', 'paragraphs': ['Русский текст 12345']})
    pdf = engine.convert_document(raw, 'letter.docx', 'pdf')
    assert pdf.startswith(b'%PDF')
    result = engine.read_document(pdf, 'letter.pdf')
    assert 'Русский текст 12345' in result['units'][0]['text']
    assert engine.preview(raw, 'letter.docx').startswith(b'\x89PNG')


def test_xlsx_pdf_conversion_and_image_pdf_conversion():
    raw = engine.create_xlsx({'sheets': [{'name': 'Report', 'rows': [['Description', 'Amount'], ['Paid', 12345]]}]})
    converted = engine.convert_document(raw, 'report.xlsx', 'pdf')
    assert converted.startswith(b'%PDF')
    assert '12345' in engine.read_document(converted, 'report.pdf')['units'][0]['text']
    im = Image.new('RGB', (500, 100), 'white')
    out = BytesIO(); im.save(out, 'PNG')
    pdf = engine.convert_document(out.getvalue(), 'photo.png', 'pdf')
    assert fitz.open(stream=pdf, filetype='pdf').page_count == 1


def test_pdf_xlsx_conversion_has_source_pages_and_explicit_best_effort_warning():
    converted = engine.convert_document(pdf_fixture(), 'source.pdf', 'xlsx')
    wb = load_workbook(BytesIO(converted), data_only=False)
    assert len(wb.sheetnames) == 4  # conversion notes + all three pages
    notes = '\n'.join(str(c.value or '') for row in wb.worksheets[0] for c in row)
    assert 'best effort' in notes.lower()
    assert any('Page 3' in str(c.value) for row in wb.worksheets[3] for c in row)
    assert any('source.pdf#page=3' in str(c.value) for row in wb.worksheets[3] for c in row)


@pytest.mark.parametrize('formula', ['=WEBSERVICE("https://example.test")', "='[remote.xlsx]Sheet1'!A1", '=HYPERLINK("file:///etc/passwd")', "=cmd|' /C calc'!A0"])
def test_explicit_active_formulas_are_rejected(formula):
    with pytest.raises(ValueError, match='(?i)(active|external)'):
        engine.create_xlsx({'sheets': [{'rows': [[{'formula': formula}]]}]})


def test_ocr_unavailable_is_explicit_not_empty_success(monkeypatch):
    # The real executable interface can fail on a deployment; isolate only that failure.
    def unavailable(*args, **kwargs):
        raise engine.pytesseract.TesseractNotFoundError()
    monkeypatch.setattr(engine.pytesseract, 'get_languages', unavailable)
    result = engine.read_document(pdf_fixture(True), 'scan.pdf', count=1)
    assert result['incomplete'] is True
    assert result['units'][0]['ocr']['incomplete'] is True
    assert result['units'][0]['ocr']['warnings']


def test_docx_embedded_picture_is_explicitly_incomplete_without_rendering():
    doc = Document(); doc.add_paragraph('Native paragraph')
    image = Image.new('RGB', (100, 50), 'white')
    out = BytesIO(); image.save(out, 'PNG')
    doc.add_picture(BytesIO(out.getvalue()))
    result = BytesIO(); doc.save(result)
    read = engine.read_document(result.getvalue(), 'with-picture.docx', count=100)
    assert read['incomplete'] is True
    assert any(unit['incomplete'] for unit in read['units'])


def test_oversized_image_rejected_before_pixel_loading(monkeypatch):
    im = Image.new('RGB', (100, 100), 'white')
    out = BytesIO(); im.save(out, 'PNG')
    monkeypatch.setattr(engine, '_PIXEL_LIMIT', 5000)
    with pytest.raises(ValueError, match='pixel'):
        engine.convert_document(out.getvalue(), 'oversized.png', 'pdf')


def test_missing_sandbox_refuses_office_conversion(monkeypatch):
    raw = engine.create_docx({'paragraphs': ['Test']})
    actual = engine.shutil.which
    monkeypatch.setattr(engine.shutil, 'which', lambda name: None if name == 'bwrap' else actual(name))
    with pytest.raises(RuntimeError, match='refusing unsandboxed'):
        engine.convert_document(raw, 'sample.docx', 'pdf')


def test_pdf_detected_table_preserves_cell_values():
    pdf = fitz.open(); page = pdf.new_page()
    for x in (50, 200, 350):
        page.draw_line((x, 50), (x, 150))
    for y in (50, 100, 150):
        page.draw_line((50, y), (350, y))
    page.insert_text((60, 80), 'Name'); page.insert_text((210, 80), 'Amount')
    page.insert_text((60, 130), 'Verified'); page.insert_text((210, 130), '34567')
    converted = engine.convert_document(pdf.tobytes(), 'table.pdf', 'xlsx')
    wb = load_workbook(BytesIO(converted))
    assert any(row[0].value == 'Verified' and row[1].value == '34567' for row in wb['Page 1'])


def test_iphone_heic_exif_preview_and_pdf():
    import pillow_heif
    im = Image.new('RGB', (640, 320), 'white')
    exif = Image.Exif(); exif[271] = 'Apple'; exif[272] = 'iPhone fixture'
    im.info['exif'] = exif.tobytes()
    heif = pillow_heif.from_pillow(im)
    out = BytesIO(); heif.save(out, quality=90)
    data = out.getvalue()
    result = engine.read_document(data, 'iphone.HEIC')
    assert result['metadata']['width'] == 640
    assert result['metadata']['exif']['Make'] == 'Apple'
    assert result['total_units'] == 1
    assert engine.preview(data, 'iphone.heif').startswith(b'\x89PNG')
    assert engine.convert_document(data, 'iphone.heic', 'pdf').startswith(b'%PDF')


def test_tiff_multipage_pagination_preview_and_conversion_preserve_all_frames():
    frames = [Image.new('RGB', (200, 100), color) for color in ('red', 'green', 'blue')]
    out = BytesIO(); frames[0].save(out, 'TIFF', save_all=True, append_images=frames[1:])
    data = out.getvalue()
    result = engine.read_document(data, 'scan.tiff', start=2, count=1)
    assert result['total_units'] == 3
    assert result['next_start'] == 3
    assert result['units'][0]['frame'] == 2
    assert result['units'][0]['source'] == 'scan.tiff#frame=2'
    third = Image.open(BytesIO(engine.preview(data, 'scan.tif', page=3)))
    assert third.getpixel((20, 20)) == (0, 0, 255)
    pdf = fitz.open(stream=engine.convert_document(data, 'scan.tiff', 'pdf'), filetype='pdf')
    assert pdf.page_count == 3
    assert engine.read_document(data, 'scan.tiff', start=3)['next_start'] is None


def test_new_xlsx_headers_centered_widths_readable_and_template_styles_unchanged():
    spec = {'sheets': [{'name': 'Объекты', 'rows': [['Наименование учреждения', 'Сумма'], ['Подробное описание объекта капитального строительства', 123]]}]}
    raw = engine.create_xlsx(spec)
    sheet = load_workbook(BytesIO(raw))['Объекты']
    assert sheet['A1'].alignment.horizontal == 'center'
    assert sheet['A1'].alignment.vertical == 'center'
    assert sheet['A1'].alignment.wrap_text is True
    assert sheet['A2'].alignment.horizontal == 'left'
    assert sheet['B2'].alignment.horizontal == 'center'
    assert sheet.column_dimensions['A'].width >= 24
    from openpyxl.styles import Alignment
    wb = Workbook(); ws = wb.active; ws.title = 'Custom'
    ws['A1'] = 'Keep'; ws['A1'].alignment = Alignment(horizontal='right')
    ws.column_dimensions['A'].width = 17
    template = BytesIO(); wb.save(template)
    updated = engine.create_xlsx({'sheets': [{'name': 'Custom', 'cells': {'B2': 123}}]}, template.getvalue())
    ws = load_workbook(BytesIO(updated))['Custom']
    assert ws['A1'].alignment.horizontal == 'right'
    assert ws.column_dimensions['A'].width == 17


def test_generated_xlsx_ooxml_rows_and_cells_are_strictly_ordered_all_sheets():
    from xml.etree import ElementTree as ET
    from openpyxl.utils.cell import coordinate_to_tuple
    raw = engine.create_xlsx({'sheets': [
        {'name': 'First', 'cells': {'XFD12': 1, 'B3': 2, 'A3': 3, 'C1': 4}},
        {'name': 'Second', 'cells': {'Z9': 5, 'A1': 6, 'B9': 7}},
    ]})
    ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(BytesIO(raw)) as archive:
        sheets = [name for name in archive.namelist() if name.startswith('xl/worksheets/sheet') and name.endswith('.xml')]
        assert len(sheets) == 2
        for name in sheets:
            rows = ET.fromstring(archive.read(name)).findall('m:sheetData/m:row', ns)
            numbers = [int(row.attrib['r']) for row in rows]
            assert numbers == sorted(set(numbers))
            for row in rows:
                coordinates = [coordinate_to_tuple(cell.attrib['r']) for cell in row.findall('m:c', ns)]
                assert all(number == int(row.attrib['r']) for number, column in coordinates)
                columns = [column for number, column in coordinates]
                assert columns == sorted(set(columns))


def test_docx_content_control_and_nested_table_text_are_not_silently_lost():
    from docx.oxml import OxmlElement
    doc = Document(); doc.add_paragraph('Before')
    sdt = OxmlElement('w:sdt'); content = OxmlElement('w:sdtContent')
    paragraph = OxmlElement('w:p'); run = OxmlElement('w:r'); text = OxmlElement('w:t')
    text.text = 'Content control visible text'
    run.append(text); paragraph.append(run); content.append(paragraph); sdt.append(content)
    doc.element.body.insert(1, sdt)
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = 'Outer table'
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = 'Nested essential value 56789'
    out = BytesIO(); doc.save(out)
    result = engine.read_document(out.getvalue(), 'controls.docx', count=100)
    combined = '\n'.join(unit['text'] for unit in result['units'])
    assert 'Content control visible text' in combined
    assert 'Nested essential value 56789' in combined
    assert result['next_start'] is None


def test_xlsx_stale_dimension_does_not_hide_actual_cells():
    wb = Workbook(); wb.active['D8'] = 'Essential value beyond declared dimension'
    out = BytesIO(); wb.save(out)
    patched = BytesIO()
    with zipfile.ZipFile(BytesIO(out.getvalue())) as source, zipfile.ZipFile(patched, 'w') as target:
        for entry in source.infolist():
            content = source.read(entry)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                content = content.replace(b'<dimension ref="D8:D8"/>', b'<dimension ref="A1:A1"/>')
            target.writestr(entry, content)
    first = engine.read_document(patched.getvalue(), 'stale.xlsx', count=3)
    assert first['total_units'] == 8
    assert first['next_start'] == 4
    final = engine.read_document(patched.getvalue(), 'stale.xlsx', start=8, count=1)
    assert final['units'][0]['cells'][0]['coordinate'] == 'D8'
    assert final['units'][0]['cells'][0]['value'] == 'Essential value beyond declared dimension'


def test_template_active_formula_rejected_even_when_unchanged():
    wb = Workbook(); wb.active['A1'] = 'https://example.invalid'
    wb.active['B1'] = '=WEBSERVICE(A1)'
    out = BytesIO(); wb.save(out)
    with pytest.raises(ValueError, match='(?i)(active|external)'):
        engine.create_xlsx({'sheets': [{'name': 'Sheet', 'cells': {'C3': 4}}]}, out.getvalue())


def test_image_formula_indirect_url_is_rejected():
    with pytest.raises(ValueError, match='(?i)(active|external)'):
        engine.create_xlsx({'sheets': [{'rows': [['https://example.invalid', {'formula': '=IMAGE(A1)'}]]}]})


def test_48mp_iphone_photo_preview_is_bounded_and_declares_scale():
    im = Image.new('RGB', (8064, 6048), 'white')
    out = BytesIO(); im.save(out, 'PNG'); im.close()
    rendered = engine.preview(out.getvalue(), 'iphone-48mp.png')
    preview = Image.open(BytesIO(rendered))
    assert preview.width * preview.height <= 20_000_000
    assert preview.info['source_width'] == '8064'
    assert preview.info['source_height'] == '6048'
    assert float(preview.info['render_scale']) < 1


def test_photo_ocr_receives_scaled_pixels_and_explicit_incomplete_flag(monkeypatch):
    # Small envelope keeps the fixture cheap; assert the real scaling path.
    monkeypatch.setattr(engine, '_RENDER_PIXEL_LIMIT', 5000, raising=False)
    im = Image.new('RGB', (200, 100), 'white')
    out = BytesIO(); im.save(out, 'PNG')
    def checked_ocr(image):
        assert image.width * image.height <= 5000
        return 'Recognized scaled text', {'attempted': True, 'languages': 'rus+eng', 'incomplete': False, 'warnings': []}
    monkeypatch.setattr(engine, '_ocr', checked_ocr)
    result = engine.read_document(out.getvalue(), 'photo.png')
    assert result['metadata']['width'] == 200
    assert result['metadata']['height'] == 100
    unit = result['units'][0]
    assert unit['render']['downsampled'] is True
    assert unit['render']['scale'] == 0.5
    assert unit['ocr']['incomplete'] is True
    assert 'scaled' in ' '.join(unit['ocr']['warnings']).lower()


def test_sandbox_marker_alone_cannot_enable_unsandboxed_office(monkeypatch):
    raw = engine.create_docx({'paragraphs': ['Synthetic marker validation']})
    monkeypatch.setenv('NEXTCLOUD_DOCUMENT_SANDBOX', '1')
    monkeypatch.setenv('HOME', '/work')
    with pytest.raises(RuntimeError, match='(?i)sandbox.*context'):
        engine.convert_document(raw, 'sample.docx', 'pdf')


def test_nested_exif_original_capture_time_and_gps_are_retained_without_inference():
    im = Image.new('RGB', (100, 100), 'white')
    exif = Image.Exif()
    exif[34665] = {36867: '2026:09:18 14:22:00', 36881: '+03:00'}
    exif[34853] = {1: 'N', 2: (55.0, 45.0, 0.0), 3: 'E', 4: (37.0, 37.0, 0.0)}
    out = BytesIO(); im.save(out, 'JPEG', exif=exif)
    metadata = engine.read_document(out.getvalue(), 'camera.jpg')['metadata']
    assert metadata['capture_time_original'] == '2026:09:18 14:22:00'
    assert metadata['capture_utc_offset_original'] == '+03:00'
    assert metadata['exif']['ExifIFD']['DateTimeOriginal'] == '2026:09:18 14:22:00'
    assert metadata['exif']['GPSInfo']['GPSLatitudeRef'] == 'N'
    json.dumps(metadata)
