from io import BytesIO
from docx import Document
from docx.shared import Cm,Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from doc_engine import create_docx

def test_official_profile_new_document():
    d=Document(BytesIO(create_docx({'profile':'official','title':'СПРАВКА','paragraphs':['Проверенный факт.'],'tables':[{'rows':[['Показатель','Значение'],['Итог','15000']]}]})))
    assert d.styles['Normal'].font.name=='Times New Roman'
    assert d.styles['Normal'].font.size==Pt(14)
    assert abs(d.sections[0].page_width-Cm(21))<1000
    assert abs(d.sections[0].left_margin-Cm(3))<1000
    assert d.paragraphs[0].alignment==WD_ALIGN_PARAGRAPH.CENTER
    assert d.paragraphs[1].alignment==WD_ALIGN_PARAGRAPH.JUSTIFY
    assert abs(d.paragraphs[1].paragraph_format.first_line_indent-Cm(1.25))<1000
    assert d.tables[0].cell(0,0).paragraphs[0].alignment==WD_ALIGN_PARAGRAPH.CENTER

def test_official_profile_preserves_template():
    t=Document();t.styles['Normal'].font.name='Arial';t.styles['Normal'].font.size=Pt(11);t.sections[0].left_margin=Cm(4)
    s=BytesIO();t.save(s)
    d=Document(BytesIO(create_docx({'profile':'official','paragraphs':['Новый текст']},s.getvalue())))
    assert d.styles['Normal'].font.name=='Arial'
    assert d.styles['Normal'].font.size==Pt(11)
    assert abs(d.sections[0].left_margin-Cm(4))<1000

def test_unrelated_word_not_forced_official():
    d=Document(BytesIO(create_docx({'title':'Личный список','paragraphs':['Купить продукты']})))
    assert d.styles['Normal'].font.name!='Times New Roman'

def test_official_heading_has_no_theme_override_or_decorative_border():
    from docx.oxml.ns import qn
    d=Document(BytesIO(create_docx({'profile':'official','title':'СПРАВКА'})))
    style=d.styles['Title'].element
    assert style.find('.//'+qn('w:pBdr')) is None
    fonts=style.find('.//'+qn('w:rFonts'))
    assert not any('theme' in key.lower() for key in fonts.attrib)
