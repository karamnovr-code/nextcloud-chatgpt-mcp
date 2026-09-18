"""Independent review regressions; synthetic inputs only."""
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
import json
import re

from docx import Document
from docx.oxml import OxmlElement
from openpyxl import Workbook
import doc_engine as engine


def test_nested_docx_content_is_extracted_or_explicitly_incomplete():
    doc = Document()
    doc.add_paragraph('Visible text')
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = 'Outer'
    nested = table.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = 'IMPORTANT_NESTED_12345'
    out = BytesIO(); doc.save(out)
    result = engine.read_document(out.getvalue(), 'nested.docx', count=100)
    assert 'IMPORTANT_NESTED_12345' in json.dumps(result) or result['incomplete']


def test_docx_content_control_is_extracted_or_explicitly_incomplete():
    doc = Document(); doc.add_paragraph('Visible text')
    control = OxmlElement('w:sdt'); content = OxmlElement('w:sdtContent')
    paragraph = OxmlElement('w:p'); run = OxmlElement('w:r'); text = OxmlElement('w:t')
    text.text = 'IMPORTANT_CONTENT_CONTROL_98765'
    run.append(text); paragraph.append(run); content.append(paragraph); control.append(content)
    doc.element.body.insert(0, control)
    out = BytesIO(); doc.save(out)
    result = engine.read_document(out.getvalue(), 'control.docx', count=100)
    assert 'IMPORTANT_CONTENT_CONTROL_98765' in json.dumps(result) or result['incomplete']


def test_xlsx_stale_dimensions_do_not_hide_real_cells():
    wb = Workbook(); wb.active['A1'] = 'Header'; wb.active['D8'] = 'IMPORTANT_REAL_ROW_888'
    raw = BytesIO(); wb.save(raw); modified = BytesIO()
    with ZipFile(BytesIO(raw.getvalue())) as source, ZipFile(modified, 'w', ZIP_DEFLATED) as dest:
        for member in source.infolist():
            data = source.read(member)
            if member.filename == 'xl/worksheets/sheet1.xml':
                data = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', data)
            dest.writestr(member, data)
    result = engine.read_document(modified.getvalue(), 'stale.xlsx', count=100)
    assert 'IMPORTANT_REAL_ROW_888' in json.dumps(result) or result['incomplete']
