"""Local, byte-oriented document processing. Never resolves document URLs.

Pagination is 1-based: PDF pages, DOCX body/header/footer blocks, XLSX physical
rows flattened in sheet order, image frames. Every read reports the exact unit
count and next_start. DOCX blocks are explicitly not physical page numbers.
OCR output always needs visual verification; incomplete is separate from that
inherent uncertainty. All output is a new byte string; callers own persistence.

Office specs:
  DOCX: title, paragraphs:[str|{text,style}], tables:[{rows:[[...]]}],
        replacements:{literal:replacement} (also across text runs).
  XLSX: sheets:[{name, rows:[[literal|{formula:'=SUM(A1:A2)'}]],
                 start_row:1, cells:{'B2':value}}].
Strings are never promoted to formulas. External/active formulas are rejected.
"""
from __future__ import annotations

from datetime import date, datetime, time
from io import BytesIO
from pathlib import Path, PurePosixPath
import math
import os
import shutil
import signal
import subprocess
import tempfile
import re
import zipfile

import pymupdf as fitz
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageOps, ExifTags, PngImagePlugin
import pytesseract
from pillow_heif import register_heif_opener

register_heif_opener()


_FORMATS = {'.pdf', '.docx', '.xlsx', '.png', '.jpg', '.jpeg', '.tif', '.tiff', '.webp', '.bmp', '.heic', '.heif'}
# These are decompression/rendering safety envelopes, not per-user usage quotas.
_ARCHIVE_EXPANSION_LIMIT = 512 * 1024 * 1024
_PIXEL_LIMIT = 80_000_000
_RENDER_PIXEL_LIMIT = 20_000_000
Image.MAX_IMAGE_PIXELS = _PIXEL_LIMIT


def _validate(data: bytes, filename: str) -> str:
    if not isinstance(data, bytes) or not data:
        raise ValueError('Document must be nonempty bytes')
    if not isinstance(filename, str) or not filename or '/' in filename or '\\' in filename or '\x00' in filename or filename in {'.', '..'}:
        raise ValueError('filename must be a basename, not a path')
    suffix = PurePosixPath(filename).suffix.lower()
    if suffix not in _FORMATS:
        raise ValueError('Unsupported document format; macros and legacy Office formats are not accepted')
    if suffix in {'.docx', '.xlsx'}:
        _check_archive(data)
    return suffix


def _check_archive(data: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 100_000 or sum(i.file_size for i in infos) > _ARCHIVE_EXPANSION_LIMIT:
                raise ValueError('Archive expansion exceeds safe memory envelope')
            names = set()
            for item in infos:
                name = item.filename
                lower = name.lower()
                if name in names or name.startswith('/') or '\\' in name or '..' in PurePosixPath(name).parts or item.flag_bits & 1:
                    raise ValueError('Unsafe archive path, duplicate member, or encryption')
                names.add(name)
                if item.file_size > 1024 * 1024 and item.file_size / max(1, item.compress_size) > 1000:
                    raise ValueError('Unsafe archive compression expansion ratio')
                if 'vbaproject' in lower or '/embeddings/' in lower or '/activex/' in lower or lower.startswith('xl/externallinks/'):
                    raise ValueError('Active content or external workbook links are not accepted')
                if lower.endswith(('.xml', '.rels')):
                    content = archive.read(item)
                    if b'<!DOCTYPE' in content.upper() or b'<!ENTITY' in content.upper():
                        raise ValueError('XML entity declarations are not accepted')
    except zipfile.BadZipFile as exc:
        raise ValueError('Invalid Office archive') from exc


def _window(start: int, count: int, total: int) -> tuple[int, int]:
    if isinstance(start, bool) or isinstance(count, bool) or not isinstance(start, int) or not isinstance(count, int) or start < 1 or count < 1:
        raise ValueError('start and count must be positive integers')
    if start > max(1, total):
        raise ValueError('start exceeds total_units')
    return start - 1, min(total, start - 1 + count)


def _result(filename, unit_type, total, begin, end, units, metadata=None, warnings=None):
    return {'filename': filename, 'unit_type': unit_type, 'total_units': total,
            'start': begin + 1, 'returned_units': len(units),
            'next_start': end + 1 if end < total else None, 'units': units,
            'metadata': metadata or {}, 'warnings': warnings or [],
            'incomplete': any(u.get('incomplete', False) for u in units)}


def _json(value):
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        return value if math.isfinite(value) else str(value)
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if isinstance(value, bytes):
        return {'byte_length': len(value), 'value_omitted': True}
    if isinstance(value, dict):
        return {str(k): _json(v) for k, v in value.items()}
    if isinstance(value, (tuple, list)):
        return [_json(v) for v in value]
    return str(value)


def _render_page(page, dpi=200):
    scale = dpi / 72
    pixels = page.rect.width * page.rect.height * scale * scale
    reduced = pixels > _RENDER_PIXEL_LIMIT
    if reduced:
        scale *= math.sqrt(_RENDER_PIXEL_LIMIT / pixels)
    pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), colorspace=fitz.csRGB, alpha=False)
    return Image.frombytes('RGB', (pix.width, pix.height), pix.samples), reduced


def _image_metadata(image):
    exif = image.getexif()
    named = {ExifTags.TAGS.get(key, str(key)): _json(value) for key, value in exif.items()}
    warnings = []
    for ifd, label, tags in ((ExifTags.IFD.Exif, 'ExifIFD', ExifTags.TAGS),
                             (ExifTags.IFD.GPSInfo, 'GPSInfo', ExifTags.GPSTAGS)):
        if ifd not in exif:
            continue
        try:
            named[label] = {tags.get(key, str(key)): _json(value) for key, value in exif.get_ifd(ifd).items()}
        except (ValueError, TypeError, KeyError, OSError, SyntaxError):
            warnings.append(f'{label} metadata could not be decoded')
    original = named.get('ExifIFD', {})
    if not isinstance(original, dict):
        original = {}
    metadata = {'width': image.width, 'height': image.height, 'format': image.format, 'exif': named,
                'capture_time_original': original.get('DateTimeOriginal', named.get('DateTimeOriginal')),
                'capture_utc_offset_original': original.get('OffsetTimeOriginal', named.get('OffsetTimeOriginal'))}
    if warnings:
        metadata['warnings'] = warnings
    return metadata


def _image_render(source):
    original_width, original_height = source.size
    if original_width * original_height > _PIXEL_LIMIT:
        raise ValueError('Image exceeds safe pixel envelope')
    oriented = ImageOps.exif_transpose(source)
    oriented_width, oriented_height = oriented.size
    scale = min(1.0, math.sqrt(_RENDER_PIXEL_LIMIT / (oriented_width * oriented_height)))
    if scale < 1:
        oriented.thumbnail((max(1, int(oriented_width * scale)), max(1, int(oriented_height * scale))),
                           Image.Resampling.LANCZOS, reducing_gap=3.0)
    image = oriented.convert('RGB')
    if image is not oriented:
        oriented.close()
    return image, {'source_width': original_width, 'source_height': original_height,
                   'width': image.width, 'height': image.height,
                   'scale': min(image.width / oriented_width, image.height / oriented_height),
                   'downsampled': scale < 1, 'orientation_applied': True}


def _ocr(image):
    info = {'attempted': True, 'languages': 'rus+eng', 'incomplete': False,
            'requires_visual_verification': True, 'warnings': []}
    try:
        languages = set(pytesseract.get_languages(config=''))
        missing = {'rus', 'eng'} - languages
        selected = [language for language in ('rus', 'eng') if language in languages]
        if missing:
            info['incomplete'] = True
            info['warnings'].append('Missing OCR languages: ' + ', '.join(sorted(missing)))
        if not selected:
            return '', info
        info['languages'] = '+'.join(selected)
        output = pytesseract.image_to_data(image, lang=info['languages'], config='--psm 3',
                                          output_type=pytesseract.Output.DICT, timeout=120)
        lines = {}
        confidence = []
        for idx, word in enumerate(output['text']):
            if not word.strip():
                continue
            key = (output['page_num'][idx], output['block_num'][idx], output['par_num'][idx], output['line_num'][idx])
            lines.setdefault(key, []).append(word)
            conf = float(output['conf'][idx])
            if conf >= 0:
                confidence.append(conf)
        text = '\n'.join(' '.join(words) for words in lines.values())
        info['mean_confidence'] = round(sum(confidence) / len(confidence), 2) if confidence else None
        if not text:
            info['incomplete'] = True
            info['warnings'].append('No text recognized; page may be blank, a photo, or unreadable')
        elif min(confidence, default=100) < 30:
            info['incomplete'] = True
            info['warnings'].append('Low confidence OCR words require visual verification')
        return text, info
    except (RuntimeError, pytesseract.TesseractNotFoundError, pytesseract.TesseractError) as exc:
        info['incomplete'] = True
        info['warnings'].append('OCR failed: ' + type(exc).__name__)
        return '', info


def _pdf_unit(page, filename):
    native = page.get_text('text', sort=True).strip()
    # OCR image-bearing pages too: a searchable heading does not make its scan readable.
    needs_ocr = bool(page.get_image_info()) or len(native) < 20 or '\ufffd' in native
    ocr = {'attempted': False, 'languages': None, 'incomplete': False, 'warnings': []}
    text = native
    if needs_ocr:
        im, reduced = _render_page(page)
        text, ocr = _ocr(im)
        if reduced:
            ocr['incomplete'] = True
            ocr['warnings'].append('Page raster resolution reduced to safe pixel envelope')
        if native:
            normalized = ' '.join(text.split()).casefold()
            supplemental = [line for line in native.splitlines() if ' '.join(line.split()).casefold() not in normalized]
            if supplemental:
                text = text + '\n\n[Native text supplement]\n' + '\n'.join(supplemental) if text else native
    return {'page': page.number + 1, 'source': f'{filename}#page={page.number + 1}',
            'text': text, 'native_text': native, 'ocr': ocr, 'incomplete': ocr['incomplete']}


def _docx_blocks(parent, element=None):
    if element is None:
        element = parent._element.body if hasattr(parent._element, 'body') else parent._element
    for child in element.iterchildren():
        if child.tag == qn('w:p'):
            yield 'paragraph', Paragraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield 'table', Table(child, parent)
        elif child.tag in {qn('w:sdt'), qn('w:sdtContent'), qn('w:customXml'), qn('w:ins')}:
            yield from _docx_blocks(parent, child)


def _paragraph_text(paragraph):
    parts = []
    for element in paragraph._element.iter():
        if element.tag == qn('w:t'):
            parts.append(element.text or '')
        elif element.tag == qn('w:tab'):
            parts.append('\t')
        elif element.tag in {qn('w:br'), qn('w:cr')}:
            parts.append('\n')
    return ''.join(parts)


def _cell_text(cell):
    parts = []
    for kind, item in _docx_blocks(cell):
        if kind == 'paragraph':
            parts.append(_paragraph_text(item))
        else:
            parts.append('\n'.join('\t'.join(_cell_text(nested) for nested in row.cells) for row in item.rows))
    return '\n'.join(parts)


def _docx_read(data, filename, start, count):
    doc = Document(BytesIO(data))
    blocks = [(f'body', kind, item) for kind, item in _docx_blocks(doc)]
    seen = set()
    for section_id, section in enumerate(doc.sections, 1):
        for name in ('header', 'footer', 'first_page_header', 'first_page_footer', 'even_page_header', 'even_page_footer'):
            container = getattr(section, name)
            if container.is_linked_to_previous or container.part.partname in seen:
                continue
            seen.add(container.part.partname)
            blocks.extend((f'section={section_id}&{name}', kind, item) for kind, item in _docx_blocks(container))
    begin, end = _window(start, count, len(blocks))
    units = []
    for index in range(begin, end):
        location, kind, item = blocks[index]
        unit = {'block': index + 1, 'kind': kind, 'source': f'{filename}#{location}&block={index+1}', 'incomplete': False}
        if kind == 'paragraph':
            unit['text'] = _paragraph_text(item)
        else:
            unit['rows'] = [[_cell_text(cell) for cell in row.cells] for row in item.rows]
            unit['text'] = '\n'.join('\t'.join(row) for row in unit['rows'])
        if item._element.xpath('.//w:drawing | .//w:pict | .//w:txbxContent'):
            unit['incomplete'] = True
            unit['warnings'] = ['Embedded picture or drawing is not extracted; render DOCX to PDF and read its pages.']
        units.append(unit)
    return _result(filename, 'block', len(blocks), begin, end, units,
                   warnings=['DOCX units are logical blocks, not rendered pages; embedded drawings and text boxes are not extracted. Use PDF rendering to inspect page layout.'])


def _xlsx_read(data, filename, start, count):
    formulas = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    values = load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        # OOXML dimension is an untrusted hint; older/exported workbooks can
        # incorrectly declare A1:A1 while containing populated cells elsewhere.
        for sheet in formulas.worksheets:
            sheet.reset_dimensions()
            sheet.calculate_dimension(force=True)
            values[sheet.title].reset_dimensions()
        totals = [(sheet.title, sheet.max_row or 0) for sheet in formulas.worksheets]
        total = sum(rows for _, rows in totals)
        begin, end = _window(start, count, total)
        units = []
        offset = 0
        for name, rows in totals:
            local_begin, local_end = max(0, begin-offset), min(rows, end-offset)
            if local_end > local_begin:
                sheet = formulas[name]
                frows = sheet.iter_rows(min_row=local_begin+1, max_row=local_end)
                vrows = values[name].iter_rows(min_row=local_begin+1, max_row=local_end, max_col=sheet.max_column)
                for row_id, (frow, vrow) in enumerate(zip(frows, vrows), local_begin+1):
                    cells = []
                    for fcell, vcell in zip(frow, vrow):
                        if fcell.value is None:
                            continue
                        is_formula = fcell.data_type == 'f'
                        cells.append({'coordinate': fcell.coordinate, 'value': _json(vcell.value if is_formula else fcell.value),
                                      'formula': _json(fcell.value) if is_formula else None,
                                      'cached_value_available': vcell.value is not None if is_formula else True,
                                      'data_type': fcell.data_type})
                    units.append({'sheet': name, 'row': row_id, 'source': f'{filename}#sheet={name}&row={row_id}',
                                  'cells': cells, 'text': '\t'.join(str(c['formula'] or c['value']) for c in cells),
                                  'incomplete': False})
            offset += rows
        return _result(filename, 'row', total, begin, end, units,
                       metadata={'sheets': [{'name': name, 'rows': rows} for name, rows in totals]},
                       warnings=['Formula values are cached values from the source, not recalculated; missing caches are explicit. Rows include physical blank rows inside sheet dimensions.'])
    finally:
        formulas.close(); values.close()


def read_document(data: bytes, filename: str, start: int = 1, count: int = 10) -> dict:
    suffix = _validate(data, filename)
    if suffix == '.pdf':
        with fitz.open(stream=data, filetype='pdf') as doc:
            if doc.needs_pass:
                raise ValueError('Encrypted PDF requires a decrypted copy')
            begin, end = _window(start, count, doc.page_count)
            units = [_pdf_unit(doc[index], filename) for index in range(begin, end)]
            return _result(filename, 'page', doc.page_count, begin, end, units, _json(doc.metadata),
                           ['OCR text requires visual verification.'] if any(u['ocr']['attempted'] for u in units) else [])
    if suffix == '.docx':
        return _docx_read(data, filename, start, count)
    if suffix == '.xlsx':
        return _xlsx_read(data, filename, start, count)
    with Image.open(BytesIO(data)) as image:
        total = getattr(image, 'n_frames', 1)
        begin, end = _window(start, count, total)
        metadata = _image_metadata(image)
        units = []
        for index in range(begin, end):
            image.seek(index)
            frame, render = _image_render(image)
            try:
                text, ocr = _ocr(frame)
            finally:
                frame.close()
            if render['downsampled']:
                ocr['incomplete'] = True
                ocr['warnings'].append('OCR used a scaled image; fine details at source resolution were not analyzed.')
            units.append({'frame': index+1, 'source': f'{filename}#frame={index+1}', 'text': text,
                          'ocr': ocr, 'render': render, 'incomplete': ocr['incomplete']})
        return _result(filename, 'frame', total, begin, end, units, metadata)


def preview(data: bytes, filename: str, page: int = 1) -> bytes:
    suffix = _validate(data, filename)
    if suffix in {'.docx', '.xlsx'}:
        return preview(convert_document(data, filename, 'pdf'), 'preview.pdf', page)
    render = None
    if suffix == '.pdf':
        with fitz.open(stream=data, filetype='pdf') as doc:
            _window(page, 1, doc.page_count)
            im, _ = _render_page(doc[page-1], dpi=120)
    else:
        with Image.open(BytesIO(data)) as source:
            _window(page, 1, getattr(source, 'n_frames', 1))
            source.seek(page-1)
            im, render = _image_render(source)
    result = BytesIO()
    info = PngImagePlugin.PngInfo()
    if render is not None:
        for key in ('source_width', 'source_height'):
            info.add_text(key, str(render[key]))
        info.add_text('render_scale', str(render['scale']))
        info.add_text('downsampled', str(render['downsampled']).lower())
    try:
        im.save(result, 'PNG', pnginfo=info)
    finally:
        im.close()
    return result.getvalue()


def _replace_paragraph(paragraph, replacements):
    # Apply literal matches across run boundaries, preserving formatting of the
    # surrounding runs. Replacement inherits formatting of its first character.
    for needle, replacement in replacements.items():
        if not isinstance(needle, str) or not needle:
            raise ValueError('Replacement keys must be nonempty literal strings')
        text = paragraph.text
        positions = list(re.finditer(re.escape(needle), text))
        for match in reversed(positions):
            cursor = 0
            for run in paragraph.runs:
                old = run.text
                left, right = cursor, cursor + len(old)
                cursor = right
                if right <= match.start() or left >= match.end():
                    continue
                prefix = old[:max(0, match.start()-left)]
                suffix = old[max(0, match.end()-left):] if match.end() < right else ''
                insert = str(replacement) if left <= match.start() < right else ''
                run.text = prefix + insert + suffix


def _replace_container(container, replacements):
    for paragraph in container.paragraphs:
        _replace_paragraph(paragraph, replacements)
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                _replace_container(cell, replacements)


def create_docx(spec: dict, template: bytes | None = None) -> bytes:
    if not isinstance(spec, dict):
        raise ValueError('spec must be an object')
    if template is not None:
        _validate(template, 'template.docx')
    doc = Document(BytesIO(template)) if template is not None else Document()
    official = template is None and spec.get('profile') == 'official'
    if official:
        for section in doc.sections:
            section.page_width, section.page_height = Cm(21), Cm(29.7)
            section.left_margin, section.right_margin = Cm(3), Cm(1.5)
            section.top_margin = section.bottom_margin = Cm(2)
        for name in ('Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3'):
            style = doc.styles[name]
            style.font.name = 'Times New Roman'
            style.font.size = Pt(14)
            style.font.color.rgb = RGBColor(0, 0, 0)
            rpr = style.element.get_or_add_rPr()
            rpr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
            rpr.rFonts.set(qn('w:cs'), 'Times New Roman')
            for attr in list(rpr.rFonts.attrib):
                if 'theme' in attr.lower():
                    del rpr.rFonts.attrib[attr]
            for tag in ('w:szCs', 'w:spacing'):
                for element in rpr.findall(qn(tag)):
                    rpr.remove(element)
            for border in style.element.findall('.//' + qn('w:pBdr')):
                border.getparent().remove(border)
            style.paragraph_format.space_before = Pt(0)
            style.paragraph_format.space_after = Pt(0)
            style.paragraph_format.line_spacing = 1.0
    replacements = spec.get('replacements', {})
    _replace_container(doc, replacements)
    for section in doc.sections:
        for name in ('header', 'footer', 'first_page_header', 'first_page_footer', 'even_page_header', 'even_page_footer'):
            container = getattr(section, name)
            if not container.is_linked_to_previous:
                _replace_container(container, replacements)
    if spec.get('title'):
        title = doc.add_heading(str(spec['title']), level=0)
        if official:
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title.paragraph_format.space_after = Pt(14)
    for item in spec.get('paragraphs', []):
        if isinstance(item, dict):
            doc.add_paragraph(str(item.get('text', '')), style=item.get('style'))
        else:
            doc.add_paragraph(str(item))
    if official:
        for paragraph in doc.paragraphs:
            if paragraph.style.name not in ('Title', 'Heading 1', 'Heading 2', 'Heading 3'):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                paragraph.paragraph_format.first_line_indent = Cm(1.25)
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for item in spec.get('tables', []):
        rows = item.get('rows', [])
        if not rows:
            continue
        table = doc.add_table(rows=0, cols=max(len(row) for row in rows))
        table.style = 'Table Grid'
        for values in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, values):
                cell.text = str(value) if value is not None else ''
    if official:
        for table in doc.tables:
            for i, row in enumerate(table.rows):
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 or len(cell.text) <= 24 else WD_ALIGN_PARAGRAPH.LEFT
                        for run in paragraph.runs:
                            run.font.name = 'Times New Roman'
                            run.font.size = Pt(14)
                            run.bold = i == 0
    out = BytesIO(); doc.save(out)
    return out.getvalue()


def _reject_active_formula(formula):
    if not isinstance(formula, str):
        raise ValueError('Unsupported formula representation')
    if re.search(r'\[|\]|\||(?:https?|file|ftp):|\b(?:WEBSERVICE|HYPERLINK|IMAGE|RTD|DDE|CALL|EXEC|EVALUATE|RUN|REGISTER(?:\.ID)?)\s*\(', formula, re.I):
        raise ValueError('External or active formula is not accepted')


def _set_cell(cell, value):
    if isinstance(value, dict):
        formula = value.get('formula')
        if set(value) != {'formula'} or not isinstance(formula, str) or not formula.startswith('='):
            raise ValueError('Explicit formulas use exactly {formula: "=..."}')
        _reject_active_formula(formula)
        cell.value = formula
    elif isinstance(value, str):
        cell.value = value
        cell.data_type = 's'
    elif value is None or isinstance(value, (int, float, bool, date, datetime, time)):
        cell.value = value
    else:
        raise ValueError('Cell value must be a scalar or explicit formula')


def _style_new_sheet(sheet):
    # Iterate only populated cells; sparse sheets can have distant coordinates.
    cells = [cell for cell in sheet._cells.values() if cell.value is not None]
    if not cells:
        return
    header_row = min(cell.row for cell in cells)
    widths = {}
    for cell in cells:
        text = str(cell.value)
        header = cell.row == header_row
        cell.alignment = Alignment(horizontal='center' if header or len(text) <= 24 else 'left',
                                   vertical='center', wrap_text=True)
        if header:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='24476A')
        width = min(60, max(12, max(len(line) for line in text.splitlines() or ['']) + 2))
        widths[cell.column] = max(widths.get(cell.column, 12), width)
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.row_dimensions[header_row].height = 32


def create_xlsx(spec: dict, template: bytes | None = None) -> bytes:
    if not isinstance(spec, dict):
        raise ValueError('spec must be an object')
    if template is not None:
        _validate(template, 'template.xlsx')
    wb = load_workbook(BytesIO(template), keep_links=False) if template is not None else Workbook()
    if template is None and spec.get('sheets'):
        wb.remove(wb.active)
    for index, item in enumerate(spec.get('sheets', []), 1):
        name = item.get('name', f'Sheet{index}')
        sheet = wb[name] if name in wb.sheetnames else wb.create_sheet(name)
        start_row = item.get('start_row', 1)
        if not isinstance(start_row, int) or start_row < 1:
            raise ValueError('start_row must be positive')
        for row_id, row in enumerate(item.get('rows', []), start_row):
            for column_id, value in enumerate(row, 1):
                _set_cell(sheet.cell(row_id, column_id), value)
        for coordinate, value in item.get('cells', {}).items():
            if not isinstance(coordinate, str) or not re.fullmatch(r'[A-Za-z]{1,3}[1-9][0-9]{0,6}', coordinate):
                raise ValueError('Invalid cell coordinate')
            row, column = coordinate_to_tuple(coordinate)
            if row > 1048576 or column > 16384:
                raise ValueError('Cell coordinate exceeds Excel dimensions')
            _set_cell(sheet.cell(row, column), value)
    # Validate retained formulas too: updating an unrelated cell must not emit
    # a template containing executable/network-capable formulas.
    for sheet in wb.worksheets:
        for cell in sheet._cells.values():
            if cell.data_type == 'f':
                _reject_active_formula(cell.value)
        for name in sheet.defined_names.values():
            _reject_active_formula(name.attr_text or '')
    for name in wb.defined_names.values():
        _reject_active_formula(name.attr_text or '')
    if template is None:
        for sheet in wb.worksheets:
            _style_new_sheet(sheet)
    out = BytesIO(); wb.save(out); wb.close()
    return out.getvalue()


def _in_worker_sandbox():
    """Trust only the fixed launcher marker plus its expected read-only mounts.

    Jobs supplies the marker in its fixed environment, never from document
    options. The entire worker is already in bwrap's network/filesystem/PID
    namespaces. Reusing it avoids unsupported nested user namespace creation.
    """
    if os.environ.get('NEXTCLOUD_DOCUMENT_SANDBOX') != '1':
        return False
    try:
        readonly_code = all(Path(path).is_file() and os.statvfs(path).f_flag & os.ST_RDONLY
                            for path in ('/app/worker.py', '/app/doc_engine.py'))
        valid = (readonly_code and Path('/work/request.json').is_file()
                 and os.getcwd() == '/app' and os.environ.get('HOME') == '/work')
    except OSError:
        valid = False
    if not valid:
        raise RuntimeError('Invalid document sandbox context; refusing Office conversion')
    return True


def _office_pdf(data: bytes, suffix: str) -> bytes:
    """Render in a filesystem + network namespace; fail closed without bubblewrap.

    Only system binaries/libraries/fonts are visible read-only. User home, host
    /tmp, network and document hyperlinks are inaccessible. A fresh LO profile
    has macros disabled. No caller-supplied filenames reach the process.
    """
    worker_sandbox = _in_worker_sandbox()
    bwrap = shutil.which('bwrap')
    if (not worker_sandbox and not bwrap) or not shutil.which('libreoffice'):
        raise RuntimeError('PDF rendering requires bubblewrap and LibreOffice; refusing unsandboxed conversion')
    with tempfile.TemporaryDirectory(prefix='document-render-', dir='/work' if worker_sandbox else None) as directory:
        root = Path(directory)
        (root / 'input' ).mkdir()
        (root / 'output').mkdir()
        (root / 'tmp').mkdir()
        (root / 'profile' / 'user').mkdir(parents=True)
        (root / 'input' / ('document' + suffix)).write_bytes(data)
        (root / 'profile' / 'user' / 'registrymodifications.xcu').write_text(
            '<?xml version="1.0" encoding="UTF-8"?>'
            '<oor:items xmlns:oor="http://openoffice.org/2001/registry">'
            '<item oor:path="/org.openoffice.Office.Common/Security/Scripting">'
            '<prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>'
            '</oor:items>', encoding='utf-8')
        if worker_sandbox:
            # The outer sandbox already exposes only this job's /work and
            # read-only runtime files. Do not create another user namespace.
            cmd = []
            render_root = root
        else:
            cmd = [bwrap, '--unshare-all', '--die-with-parent', '--new-session',
                   '--ro-bind', '/usr', '/usr', '--ro-bind', '/lib', '/lib',
                   '--ro-bind', '/lib64', '/lib64', '--symlink', 'usr/bin', '/bin',
                   '--symlink', 'usr/sbin', '/sbin', '--proc', '/proc', '--dev', '/dev',
                   '--tmpfs', '/tmp', '--bind', directory, '/work', '--chdir', '/work']
            for system_path in ('/etc/fonts', '/etc/ld.so.cache', '/etc/libreoffice'):
                if Path(system_path).exists():
                    cmd.extend(['--ro-bind', system_path, system_path])
            render_root = Path('/work')
        cmd.extend(['/usr/bin/libreoffice', '-env:UserInstallation=' + (render_root / 'profile').as_uri(),
                    '--headless', '--nologo', '--nodefault', '--nofirststartwizard',
                    '--convert-to', 'pdf', '--outdir', str(render_root / 'output'),
                    str(render_root / 'input' / ('document' + suffix))])
        env = {'PATH': '/usr/bin:/bin', 'HOME': str(render_root), 'TMPDIR': str(render_root / 'tmp'),
               'XDG_CACHE_HOME': str(render_root / 'cache'), 'LANG': 'C.UTF-8', 'SAL_USE_VCLPLUGIN': 'svp',
               'OMP_THREAD_LIMIT': '2'}
        process = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                                   env=env, cwd=directory, start_new_session=True)
        try:
            stdout, stderr = process.communicate(timeout=180)
        except subprocess.TimeoutExpired:
            os.killpg(process.pid, signal.SIGKILL)
            process.communicate()
            raise RuntimeError('Sandboxed Office rendering timed out') from None
        output = root / 'output' / 'document.pdf'
        if process.returncode or not output.is_file():
            # Do not surface arbitrary document-derived process output as instructions.
            raise RuntimeError(f'Sandboxed Office rendering failed (exit={process.returncode})')
        result = output.read_bytes()
        if not result.startswith(b'%PDF'):
            raise RuntimeError('Office renderer did not produce a PDF')
        return result


def _pdf_xlsx(data, filename):
    sheets = [{'name': 'Conversion notes', 'rows': [
        ['Best effort table extraction; verify all numbers and structure against the PDF.'],
        ['Image-only tables use OCR text rows; columns and merged cells are not reconstructed.'],
        ['No detected table: extracted page text is preserved as rows.'],
        ['Source', filename]]}]
    with fitz.open(stream=data, filetype='pdf') as pdf:
        if pdf.needs_pass:
            raise ValueError('Encrypted PDF requires a decrypted copy')
        for page in pdf:
            unit = _pdf_unit(page, filename)
            rows = [['Source', unit['source']]]
            rows.append(['OCR attempted', unit['ocr']['attempted'], 'OCR incomplete', unit['incomplete']])
            tables = page.find_tables().tables
            if tables:
                for number, table in enumerate(tables, 1):
                    rows.append([f'Table {number}: best effort, verify against page'])
                    rows.extend(table.extract())
                    rows.append([])
            # Retain text even when table detection finds only a partial table.
            rows.append(['Full extracted page text (verify OCR when present)'])
            rows.extend([[line] for line in unit['text'].splitlines()])
            sheets.append({'name': f'Page {page.number + 1}', 'rows': rows})
    return create_xlsx({'sheets': sheets})


def convert_document(data: bytes, filename: str, target_format: str) -> bytes:
    suffix = _validate(data, filename)
    target = target_format.lower().lstrip('.')
    if target == suffix.lstrip('.'):
        return data
    if target == 'pdf' and suffix in {'.docx', '.xlsx'}:
        return _office_pdf(data, suffix)
    if target == 'pdf' and suffix not in {'.pdf', '.docx', '.xlsx'}:
        with Image.open(BytesIO(data)) as image:
            frames = []
            for index in range(getattr(image, 'n_frames', 1)):
                image.seek(index)
                if image.width * image.height > _PIXEL_LIMIT:
                    raise ValueError('Image exceeds safe pixel envelope')
                frames.append(ImageOps.exif_transpose(image).convert('RGB'))
            out = BytesIO()
            frames[0].save(out, 'PDF', save_all=True, append_images=frames[1:])
            return out.getvalue()
    if suffix == '.pdf' and target == 'xlsx':
        return _pdf_xlsx(data, filename)
    if suffix == '.pdf' and target == 'docx':
        doc = Document()
        doc.core_properties.comments = 'Best effort text extraction; layout and tables are not reconstructed; OCR may be incomplete.'
        with fitz.open(stream=data, filetype='pdf') as pdf:
            for index, page in enumerate(pdf):
                if index:
                    doc.add_page_break()
                unit = _pdf_unit(page, filename)
                doc.add_paragraph(f'Источник: {unit["source"]}')
                if unit['ocr']['attempted']:
                    doc.add_paragraph('Распознано OCR; требуется визуальная проверка.' + (' Распознавание неполное.' if unit['incomplete'] else ''))
                doc.add_paragraph(unit['text'])
        out = BytesIO(); doc.save(out)
        return out.getvalue()
    raise ValueError('Unsupported conversion')
